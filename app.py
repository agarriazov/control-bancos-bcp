import streamlit as st
import pandas as pd
import openpyxl
import io

st.set_page_config(page_title="Control de Ingresos BCP", page_icon="🏦")
st.title("🏦 Control de Ingresos Bancarios BCP")
st.markdown("Carga tu archivo maestro y los reportes diarios o históricos para actualizar la base automáticamente.")

def limpiar_op_para_clave(val):
    if pd.isna(val) or val is None:
        return ""
    s = str(val).split('.')[0].strip()
    return s.lstrip('0')

def formatear_op_visual(val):
    if pd.isna(val) or val is None:
        return ""
    return str(val).split('.')[0].strip()

def clave_unica(cuenta, num_op, monto):
    c = str(cuenta).strip()
    o = limpiar_op_para_clave(num_op)
    try:
        m = f"{float(monto):.2f}"
    except:
        m = str(monto).strip()
    return f"{c}_{o}_{m}"

st.subheader("1. Subir Archivos")
archivo_maestro = st.file_uploader("Sube tu archivo actual (BCP Diario 2026.xlsm):", type=["xlsm", "xlsx"])
archivos_insumos = st.file_uploader("Sube los reportes descargados del BCP (puedes seleccionar varios):", type=["xlsx", "xls"], accept_multiple_files=True)

if st.button("🚀 Procesar y Actualizar Ingresos", type="primary"):
    if not archivo_maestro:
        st.error("Debes subir primero el archivo Maestro (BCP Diario 2026.xlsm).")
    elif not archivos_insumos:
        st.warning("Debes subir al menos un reporte del BCP.")
    else:
        with st.spinner("Procesando datos..."):
            wb = openpyxl.load_workbook(archivo_maestro, keep_vba=True)
            ws = wb['Movimientos']

            # Limpieza de egresos históricos (monto <= 0)
            filas_eliminadas = 0
            for r in range(ws.max_row, 2, -1):
                m = ws.cell(r, 5).value
                if m is not None:
                    try:
                        if float(m) <= 0:
                            ws.delete_rows(r, 1)
                            filas_eliminadas += 1
                    except:
                        pass

            # Mapeo de ingresos existentes
            registros_existentes = {}
            for r in range(3, ws.max_row + 1):
                cuenta = ws.cell(r, 9).value
                num_op = ws.cell(r, 7).value
                monto = ws.cell(r, 5).value
                if cuenta and num_op and monto is not None:
                    k = clave_unica(cuenta, num_op, monto)
                    if k not in registros_existentes:
                        registros_existentes[k] = r

            nuevos = 0
            actualizados = 0

            for f in archivos_insumos:
                xl = pd.ExcelFile(f)
                for sheet in xl.sheet_names:
                    df_raw = pd.read_excel(f, sheet_name=sheet, header=None)
                    if len(df_raw) < 5:
                        continue
                    cuenta = str(df_raw.iloc[0, 1]).strip()
                    moneda = str(df_raw.iloc[1, 1]).strip()
                    cols = [str(c).strip() for c in df_raw.iloc[4].tolist()]
                    df_data = pd.read_excel(f, sheet_name=sheet, skiprows=5, header=None)

                    es_hist = any('Operación - Número' in c for c in cols)

                    for _, fila in df_data.iterrows():
                        try:
                            monto = float(fila[3])
                        except:
                            continue

                        if monto <= 0:
                            continue

                        fecha_val = pd.to_datetime(fila[0], errors='coerce', dayfirst=True)
                        desc = str(fila[2]).strip()
                        sucursal = str(fila[4]).strip()
                        num_op_raw = fila[5]
                        usuario = str(fila[7]).strip() if es_hist and len(fila) > 7 else (str(fila[6]).strip() if len(fila) > 6 else '')

                        k = clave_unica(cuenta, num_op_raw, monto)

                        if k in registros_existentes:
                            if es_hist:
                                fila_dest = registros_existentes[k]
                                if pd.notna(fecha_val):
                                    ws.cell(fila_dest, 2).value = fecha_val.strftime('%d/%m/%Y')
                                ws.cell(fila_dest, 4).value = desc
                                ws.cell(fila_dest, 8).value = usuario
                                actualizados += 1
                        else:
                            nueva_pos = ws.max_row + 1
                            registros_existentes[k] = nueva_pos

                            concat_str = f"{fecha_val.strftime('%d/%m/%Y') if pd.notna(fecha_val) else ''}-{cuenta}-{formatear_op_visual(num_op_raw)}-{sucursal}-{desc}-"

                            ws.cell(nueva_pos, 1).value = concat_str
                            ws.cell(nueva_pos, 2).value = fecha_val.strftime('%d/%m/%Y') if pd.notna(fecha_val) else ''
                            ws.cell(nueva_pos, 3).value = None
                            ws.cell(nueva_pos, 4).value = desc
                            ws.cell(nueva_pos, 5).value = monto
                            ws.cell(nueva_pos, 6).value = sucursal

                            cell_op = ws.cell(nueva_pos, 7)
                            cell_op.value = formatear_op_visual(num_op_raw)
                            cell_op.number_format = '@'

                            ws.cell(nueva_pos, 8).value = usuario
                            ws.cell(nueva_pos, 9).value = cuenta
                            ws.cell(nueva_pos, 10).value = moneda
                            ws.cell(nueva_pos, 11).value = ""
                            nuevos += 1

            for col_oculta in ['C', 'F', 'H', 'I']:
                ws.column_dimensions[col_oculta].hidden = True

            output_buffer = io.BytesIO()
            wb.save(output_buffer)
            output_buffer.seek(0)

            st.success("✅ ¡Procesamiento completado con éxito!")
            st.write(f"- Operaciones actualizadas con histórico: **{actualizados}**")
            st.write(f"- Nuevos ingresos agregados: **{nuevos}**")

            st.download_button(
                label="📥 Descargar BCP Diario 2026.xlsm Actualizado",
                data=output_buffer,
                file_name="BCP Diario 2026.xlsm",
                mime="application/vnd.ms-excel.sheet.macroEnabled.12"
            )
